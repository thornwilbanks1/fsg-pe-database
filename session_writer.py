"""
Session writer — called by Claude Code to append validated portcos to the output xlsx.
Usage: python3 session_writer.py '<json_array>'
"""
import sys, json, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from rapidfuzz import fuzz

OUTPUT_DB = "STEPHENS_FSG_DATABASE.xlsx"

SCHEMA_COLS = [
    "pe_firm", "portfolio_company", "status", "sector", "sub_sector",
    "hq", "transaction_type", "investing_fund", "entry_date",
    "prior_owner", "sell_side_advisor", "entry_press_release",
    "pe_firm_deal_team", "exit_date", "exit_buyer", "buyer_type",
    "exit_press_release", "website", "company_description",
    "business_characteristics", "end_markets", "data_quality_notes",
]

VALID_STATUS     = {"Current", "Realized"}
VALID_BUYER_TYPE = {"Sponsor", "Strategic", "IPO", "Other", "---", ""}
VALID_TX_TYPE    = {
    "Buyout", "Growth Capital (Minority)", "Growth Equity",
    "Co-Investment", "Carve-Out", "Carve-Out (from Corporate)",
    "Public-to-Private", "Recap", "",
}
DATE_RE = re.compile(r"^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)-\d{2}$")
FUZZY_THRESH = 88

GREEN = PatternFill("solid", fgColor="E8F5E9")
BLUE  = PatternFill("solid", fgColor="D6E4F0")
WRAP  = Alignment(wrap_text=True, vertical="top")
THIN  = Side(style="thin", color="CCCCCC")
BDR   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DFONT = Font(name="Arial", size=9)

NULL_CURRENT = {"exit_date", "exit_buyer", "buyer_type", "exit_press_release"}


def load_existing(ws):
    existing = set()
    for r in range(2, ws.max_row + 1):
        f = str(ws.cell(r, 1).value or "").strip().lower()
        c = str(ws.cell(r, 2).value or "").strip().lower()
        if f and c:
            existing.add((f, c))
    return existing


def normalize(row):
    # Apply null conventions
    if row.get("status") == "Current":
        for col in NULL_CURRENT:
            row[col] = "---"
    # Clean website
    site = row.get("website", "")
    if site:
        site = re.sub(r"^https?://", "", site).rstrip("/")
        site = re.sub(r"^www\.", "", site).split("/")[0]
        row["website"] = site
    # Default nulls
    for col in SCHEMA_COLS:
        if col not in row or row[col] is None or row[col] == "":
            if col in NULL_CURRENT and row.get("status") == "Current":
                row[col] = "---"
            elif col == "data_quality_notes":
                row[col] = "---"
            else:
                row[col] = "Not found publicly available"
    return row


def validate(row, existing):
    hard, warn = [], []
    for f in ("pe_firm", "portfolio_company", "status"):
        if not str(row.get(f, "")).strip():
            hard.append(f"R01: '{f}' blank")
    if row.get("status") not in VALID_STATUS:
        hard.append(f"R02: status='{row.get('status')}' invalid")
    bt = row.get("buyer_type", "")
    if bt not in VALID_BUYER_TYPE:
        hard.append(f"R03: buyer_type='{bt}' invalid")
    for df in ("entry_date", "exit_date"):
        val = row.get(df, "")
        if val and val not in ("---", "", "Not found publicly available"):
            if not DATE_RE.match(val):
                hard.append(f"R05: {df}='{val}' not Mon-YY")
    key = (row.get("pe_firm","").lower(), row.get("portfolio_company","").lower())
    if key in existing:
        hard.append(f"R07: duplicate")
    firm_l = row.get("pe_firm","").lower()
    portco_l = row.get("portfolio_company","").lower()
    for (ef, ec) in existing:
        if ef == firm_l and ec != portco_l:
            if fuzz.ratio(ec, portco_l) >= FUZZY_THRESH:
                warn.append(f"W06: near-dup '{ec}'")
    return len(hard) == 0, hard, warn


def write_row(ws, row, existing):
    last = ws.max_row + 1
    fill = GREEN if row.get("status") == "Current" else BLUE
    for ci, col in enumerate(SCHEMA_COLS, 1):
        val = row.get(col, "")
        cell = ws.cell(last, ci, val)
        cell.fill = fill
        cell.alignment = WRAP
        cell.font = DFONT
    ws.row_dimensions[last].height = 55
    existing.add((row["pe_firm"].lower(), row["portfolio_company"].lower()))


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 session_writer.py '<json>'")
        sys.exit(1)

    portcos = json.loads(sys.argv[1])
    if not isinstance(portcos, list):
        portcos = [portcos]

    wb = openpyxl.load_workbook(OUTPUT_DB)
    ws = wb["Investments"]
    existing = load_existing(ws)

    written = skipped = failed = 0
    for p in portcos:
        p = normalize(p)
        ok, hard, warn = validate(p, existing)
        if warn:
            notes = p.get("data_quality_notes", "---")
            extra = "; ".join(warn)
            p["data_quality_notes"] = (f"{notes} | WARN: {extra}".strip(" |")
                                        if notes != "---" else f"WARN: {extra}")
        if ok:
            write_row(ws, p, existing)
            written += 1
        else:
            print(f"  FAIL {p.get('portfolio_company')}: {hard}", file=sys.stderr)
            failed += 1

    wb.save(OUTPUT_DB)
    print(f"written={written} failed={failed}")


if __name__ == "__main__":
    main()
